"""產生 downloads/ 裡的固定版本檔案：
- 學習單_3-1 ~ 3-4.pdf（worksheets.py）
- 複習卷A / 複習卷B：.pdf 與 .docx（從題庫挑「沒有圖、難度 1–2」的題目，A、B 不重複；附解答詳解）
並產生 ch3/js/04_downloads.js 給網頁列出下載連結。
用法：python3 make_downloads.py（需要 Google Chrome 與 python-docx）"""
import json, random, re, subprocess, pathlib, html as H
from worksheets import WS

root = pathlib.Path(__file__).parent
OUT = root.parent / 'downloads'; OUT.mkdir(exist_ok=True)
TMP = root.parent / 'research' / '_tmp_print'; TMP.mkdir(parents=True, exist_ok=True)
CHROME = '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'
L = 'ABCD'

CSS = '''@page { size: A4; margin: 14mm 14mm 16mm; }
body { font-family: "PingFang TC", "Heiti TC", sans-serif; font-size: 11.5pt; line-height: 1.6; color: #000; }
h1 { font-size: 16pt; margin: 0; } .sub { color: #444; font-size: 10pt; }
.head { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2px solid #000; padding-bottom: 4px; margin-bottom: 6px; }
.info { margin: 6px 0 10px; }
h2 { font-size: 12.5pt; background: #EEE; padding: 3px 8px; margin: 14px 0 6px; }
table.t { border-collapse: collapse; width: 100%; margin: 4px 0 6px; font-size: 10.5pt; }
table.t th, table.t td { border: 1px solid #555; padding: 4px 6px; text-align: center; height: 22px; }
table.t th { background: #F2F2F2; }
.blank { display: inline-block; min-width: 5em; border-bottom: 1px solid #000; margin: 0 3px; }
.line { border-bottom: 1px solid #000; height: 26px; }
.q { margin: 6px 0 10px; break-inside: avoid; } .n { font-weight: bold; }
.opts { display: grid; grid-template-columns: 1fr 1fr; gap: 0 18px; margin-left: 1.6em; }
.num { margin-left: 1.6em; }
.stem { font-weight: bold; margin-top: 10px; }
.pb { break-before: page; }
.a { margin: 3px 0; font-size: 10.5pt; break-inside: avoid; }
.foot { margin-top: 14px; font-size: 9pt; color: #666; }'''

def pdf(name, body, title):
    f = TMP / (name + '.html')
    f.write_text(f'<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><title>{title}</title><style>{CSS}</style></head><body>{body}</body></html>', encoding='utf-8')
    out = OUT / (name + '.pdf')
    subprocess.run([CHROME, '--headless=new', '--disable-gpu', '--no-pdf-header-footer', f'--print-to-pdf={out}', f.as_uri()], check=True, capture_output=True)
    return out

def head(title, sub):
    return f'<div class="head"><div><h1>{title}</h1><div class="sub">{sub}</div></div><div class="sub">勝利國中 國三理化</div></div><div class="info">班級：＿＿＿＿　座號：＿＿＿　姓名：＿＿＿＿＿＿</div>'

files = []
# ---------- 學習單 ----------
for key, (title, tabs, secs) in WS.items():
    body = head(f'學習單｜{title}', f'搭配「第三章 功與能 互動課本」{tabs}的模擬，邊操作邊記錄') + ''.join(f'<h2>{h}</h2>{c}' for h, c in secs)
    body += '<div class="foot">說明：表格中的數值請操作網頁模擬後填入。g 值請依老師指定（網頁右上角可切換 9.8 或 10）。</div>'
    pdf(f'學習單_{key}', body, title)
    files.append({'t': f'學習單 {title}', 'href': f'downloads/學習單_{key}.pdf', 'fmt': 'PDF', 'note': f'搭配網頁{tabs}的模擬操作與紀錄。'})

# ---------- 複習卷 A／B ----------
Q = json.loads((root / 'ch3' / 'js' / '03_qbank.js').read_text(encoding='utf-8').split('const QBANK = ', 1)[1].rstrip().rstrip(';'))
single = [q for q in Q if q['type'] in ('mc', 'num') and not q.get('fig') and q['lv'] in (1, 2)]
groups = [q for q in Q if q['type'] == 'group' and not q.get('fig') and not any(s.get('fig') for s in q['sub']) and q['lv'] in (1, 2)]
QUOTA = {'3-1': 6, '3-2': 7, '3-3': 6, '3-4': 7}
rng = random.Random(20261013)
pick = {'A': [], 'B': []}
for sec, k in QUOTA.items():
    cand = [q for q in single if q['sec'] == sec]; rng.shuffle(cand)
    pick['A'] += cand[:k]; pick['B'] += cand[k:2 * k]
gs = groups[:]; rng.shuffle(gs)
gA, gB = gs[:1], gs[1:2]

def clean(s):  # 給 Word 用：去掉 HTML 標籤
    s = re.sub(r'<br\s*/?>', '\n', s); s = re.sub(r'<sub>(.*?)</sub>', r'\1', s); s = re.sub(r'<sup>2</sup>', '²', s); s = re.sub(r'<sup>(.*?)</sup>', r'^\1', s)
    return H.unescape(re.sub(r'<[^>]+>', '', s))

def paper(name, qs, grp):
    qs = sorted(qs, key=lambda q: (q['sec'], q['lv']))
    seq = [('q', q) for q in qs] + [('g', g) for g in grp]
    n, body, ans = 0, '', []
    for kind, q in seq:
        if kind == 'g':
            body += f'<div class="stem">◎ 題組（第 {n + 1}～{n + len(q["sub"])} 題）：{q["stem"]}</div>'
            subs = q['sub']
        else: subs = [q]
        for s in subs:
            n += 1; ans.append((n, s))
            o = f'<div class="num">答：＿＿＿＿＿＿ {s.get("unit", "")}</div>' if s.get('type') == 'num' else '<div class="opts">' + ''.join(f'<span>({L[i]}) {x}</span>' for i, x in enumerate(s['opts'])) + '</div>'
            body += f'<div class="q"><span class="n">{n}.</span> {s["q"]}{o}</div>'
    a = ''.join(f'<div class="a"><b>{k}. {s["ans"] if s.get("type") == "num" else "(" + L[s["ans"]] + ")"}{" " + s.get("unit", "") if s.get("type") == "num" else ""}</b>　{s["exp"]}　<span class="sub">〔{s.get("src", "原創")}〕</span></div>' for k, s in ans)
    full = head(f'第 3 章 功與能　複習卷 {name}', f'範圍：3-1～3-4　共 {n} 題') + body + f'<div class="pb"></div><h1>複習卷 {name}　解答與詳解</h1>' + a
    pdf(f'複習卷{name}', full, f'複習卷{name}')
    # Word
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    d = Document()
    st = d.styles['Normal']; st.font.name = 'Microsoft JhengHei'; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    d.add_heading(f'國三理化　第 3 章 功與能　複習卷 {name}', 1)
    d.add_paragraph('班級：＿＿＿＿　座號：＿＿＿　姓名：＿＿＿＿＿＿　得分：＿＿＿＿')
    k = 0
    for kind, q in seq:
        if kind == 'g':
            p = d.add_paragraph(); p.add_run(f'◎ 題組（第 {k + 1}～{k + len(q["sub"])} 題）：{clean(q["stem"])}').bold = True
            subs = q['sub']
        else: subs = [q]
        for s in subs:
            k += 1
            d.add_paragraph(f'{k}. {clean(s["q"])}')
            if s.get('type') == 'num': d.add_paragraph(f'    答：＿＿＿＿＿＿ {s.get("unit", "")}')
            else: d.add_paragraph('    ' + '　'.join(f'({L[i]}) {clean(x)}' for i, x in enumerate(s['opts'])))
    d.add_page_break(); d.add_heading(f'複習卷 {name}　解答與詳解', 1)
    for kk, s in ans:
        p = d.add_paragraph(); r = p.add_run(f'{kk}. ' + (f'{s["ans"]} {s.get("unit", "")}' if s.get('type') == 'num' else f'({L[s["ans"]]})')); r.bold = True
        p.add_run('　' + clean(s['exp']) + f'〔{s.get("src", "原創")}〕')
    d.save(OUT / f'複習卷{name}.docx')
    return n

for name, qs, g in [('A', pick['A'], gA), ('B', pick['B'], gB)]:
    n = paper(name, qs, g)
    files.append({'t': f'複習卷 {name}（{n} 題，附解答詳解）', 'href': f'downloads/複習卷{name}.pdf', 'fmt': 'PDF', 'note': '段考難度，第 3 章全範圍。最後一頁是解答與詳解。'})
    files.append({'t': f'複習卷 {name}（Word 可編輯版）', 'href': f'downloads/複習卷{name}.docx', 'fmt': 'Word', 'note': '老師可自行刪改題目後再印。'})

(root / 'ch3' / 'js' / '04_downloads.js').write_text('/* 自動產生（make_downloads.py） */\nconst DOWNLOADS = ' + json.dumps(files, ensure_ascii=False, indent=1) + ';\n', encoding='utf-8')
print('\n'.join(f['href'] for f in files))
